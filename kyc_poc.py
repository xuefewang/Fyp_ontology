#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
IS4920 本体 PoC —— 输入案例 -> RiskFactorHit + RiskAssessmentResult

严格按概念模型流程图：
    输入实体   party1（法人）/ party2（股东链，UBO 在其中）/ Document / Application
               / Industry（industry_tier 为输入属性）/ WatchList（listed-in 关系）
    输出实体   RiskFactorHit / RiskAssessmentResult
    不做       risk appetite / escalation / mitigation / compliance decision / EDD

EDD 触发逻辑不在本次本体范围内，故 RiskFactorHit.edd_triggered 与
RiskAssessmentResult.edd_required 字段保留（流程图中有），但恒为 False。

输入：test_cases_data.json（39 个案例）。规则目录（RF-01..10 的动作类型与分值、
风险带区间）从 SSOT 工作簿实时解析，不硬编码。

运行：
    python kyc_poc.py                     # 汇总 39 个案例
    python kyc_poc.py --detail            # 逐案例打印两张结果表
    python kyc_poc.py --detail case11     # 只看单个案例
    python kyc_poc.py --builtin           # 跑内置 15 个案例（不依赖 json）
"""

from __future__ import annotations

import argparse
import datetime as dt
import itertools
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

SSOT_XLSX = Path(r"C:\Users\hfyud\Downloads\IS4920_first_phase_v2.0.xlsx")
DEFAULT_JSON = Path(r"C:\Users\hfyud\Downloads\test_cases_data.json")

# ======================================================================================
# 1. 输入实体（对应流程图的 6 类输入实体及其字段）
# ======================================================================================


@dataclass
class WatchList:
    """案例自带的名单条目，表达 listed-in 关系（该主体被列入了哪些名单）。"""

    watchlist_id: str
    list_name: str
    list_version_date: str = ""


@dataclass
class Industry:
    industry_id: str
    industry_name: str
    industry_tier: str  # I1..I4，输入直接给定


@dataclass
class Application:
    application_id: str
    submission_date: str
    application_status: str = "Submitted"


@dataclass
class Document:
    document_id: str
    document_type: str
    status: str  # Provided / Missing / Expired / Conflicting
    fact_verified: str
    is_mandatory: bool


@dataclass
class Party:
    """party1 / party2 共用结构。party1 为法人（多 full_name），party2 为股东链上的当事方。"""

    party_id: str
    person_id: str = ""
    full_name: str = ""
    is_ubo: bool = False
    is_nominee: bool = False
    is_trust: bool = False
    sanctions_flag: bool = False
    shareholding_percentage: float = 0.0
    pep_status: str = "None"
    nationality: str = ""
    registration_jurisdiction: str = ""


@dataclass
class Case:
    name: str
    application: Application
    industry: Industry
    party1: Party
    party2_chain: list[Party] = field(default_factory=list)  # 股东链，按层级顺序
    documents: list[Document] = field(default_factory=list)
    watchlists: list[WatchList] = field(default_factory=list)
    assessment_date: str = ""

    def all_parties(self) -> list[Party]:
        return [self.party1] + self.party2_chain

    def ubos(self) -> list[Party]:
        return [p for p in self.all_parties() if p.is_ubo]

    def ownership_layers(self) -> int:
        """股权链层数：party1 + party2 链上各方，共 1 + len(party2_chain) 层。"""
        return 1 + len(self.party2_chain)


# ======================================================================================
# 2. 输出实体（对应流程图右侧两张结果表，字段一一对应）
# ======================================================================================

HARD_STOP = "HardStop"
SCORE = "Score"
RETURN = "ReturnForSupplementation"
ESCALATE = "EscalateToHumanReview"


@dataclass
class RiskFactorHit:
    hit_id: str
    factor_id: str
    points_applied: int
    action_type: str
    has_hard_stop: bool
    edd_triggered: bool
    evidence: str = ""


@dataclass
class RiskAssessmentResult:
    assessment_id: str
    total_score: int
    risk_band: str
    has_hard_stop: bool
    assessment_date: str
    assessment_status: str
    edd_required: bool


# ======================================================================================
# 3. 规则目录：从 SSOT 解析
# ======================================================================================


@dataclass
class FactorRule:
    factor_id: str
    dimension: str
    name: str
    criteria: str
    action_type: str
    points: int


@dataclass
class BandRule:
    band: str
    low: int
    high: int | None


@dataclass
class RuleBook:
    factors: dict[str, FactorRule]
    bands: list[BandRule]


def _parse_action(raw: str) -> tuple[str, int]:
    s = (raw or "").lower()
    if "hard stop" in s:
        return HARD_STOP, 0
    if "return for supplementation" in s:
        return RETURN, 0
    if "escalate to human review" in s:
        return ESCALATE, 0
    m = re.search(r"\+(\d+)\s*points", s)
    return SCORE, int(m.group(1)) if m else 0


def _parse_range(raw: str) -> tuple[int, int | None]:
    s = (raw or "").replace("–", "-").replace("—", "-").lower()
    m = re.search(r"(\d+)\s*-\s*(\d+)", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.search(r"(\d+)\s*points\s*and\s*above", s)
    if m:
        return int(m.group(1)), None
    m = re.search(r"(\d+)", s)
    return (int(m.group(1)), int(m.group(1))) if m else (0, 0)


def load_rules(xlsx: Path = SSOT_XLSX) -> RuleBook:
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        raise SystemExit("需要 openpyxl：pip install openpyxl")

    wb = openpyxl.load_workbook(xlsx, data_only=True)

    factors: dict[str, FactorRule] = {}
    for row in wb["01_Risk_Factors"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        fid = str(row[0]).strip()
        atype, pts = _parse_action(str(row[4]))
        factors[fid] = FactorRule(fid, str(row[1] or ""), str(row[2] or ""), str(row[3] or ""), atype, pts)

    bands: list[BandRule] = []
    for row in wb["04_Risk_Bands"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or str(row[0]).strip().lower().startswith("special"):
            continue  # 硬停行不参与分值区间
        low, high = _parse_range(str(row[1] or ""))
        bands.append(BandRule(str(row[0]).strip(), low, high))
    bands.sort(key=lambda b: b.low)

    return RuleBook(factors=factors, bands=bands)


# ======================================================================================
# 4. Stage 1：风险因子检测 -> RiskFactorHit
#    每条判定谓词后注明其 Determination Criteria（01_Risk_Factors 第 4 列）
# ======================================================================================

BLACK, GREY, UN = "FATF黑名单", "FATF灰名单", "UN制裁名单"

# RF-08 Criteria：Crypto assets, gambling, weapons, precious metals
HIGH_RISK_INDUSTRY = ("crypto", "virtual asset", "vasp", "gambl", "casino",
                      "weapon", "defence", "defense", "dual-use", "precious metal")


def list_kind(wl: WatchList) -> str | None:
    """把案例自带的名单条目归类为 FATF 黑名单 / 灰名单 / UN 制裁名单。

    PEP 指引（HKMA PEP Guideline）属判定依据来源而非命中名单，故返回 None。
    """
    parts = wl.watchlist_id.upper().split("-")
    name = wl.list_name.lower()
    if "BL" in parts or "blacklist" in name or "high-risk jurisdictions" in name:
        return BLACK
    if "GL" in parts or "grey" in name or "increased monitoring" in name:
        return GREY
    if "UN" in parts or "united nations" in name or "security council" in name or "consolidated list" in name:
        return UN
    return None


def detect(case: Case, rb: RuleBook) -> list[RiskFactorHit]:
    hits: list[RiskFactorHit] = []
    seq = itertools.count(1)

    def add(fid: str, evidence: str) -> None:
        r = rb.factors[fid]
        hs = r.action_type == HARD_STOP
        hits.append(
            RiskFactorHit(
                hit_id=f"HIT-{next(seq):03d}",
                factor_id=fid,
                points_applied=0 if hs else r.points,  # 硬停不计分
                action_type=r.action_type,
                has_hard_stop=hs,
                edd_triggered=False,  # EDD 不在本次本体范围
                evidence=evidence,
            )
        )

    kinds = {k for k in (list_kind(w) for w in case.watchlists) if k}
    declared = ", ".join(f"{w.watchlist_id}({list_kind(w)})" for w in case.watchlists) or "无"
    ubos = case.ubos()
    ubo_txt = ", ".join(f"{p.party_id}[nat={p.nationality},jur={p.registration_jurisdiction}]" for p in ubos) or "无"

    # RF-01 地理风险：Registration or UBO nationality is North Korea, Iran, or Myanmar
    if BLACK in kinds:
        add("RF-01", f"主体被列入 FATF 黑名单；UBO：{ubo_txt}；名单：{declared}")
    # RF-02 地理风险：Registration or UBO nationality is on FATF increased monitoring list
    elif GREY in kinds:  # 与 RF-01 互斥，硬停优先
        add("RF-02", f"主体被列入 FATF 灰名单；UBO：{ubo_txt}；名单：{declared}")

    # RF-03 个人风险：Match with UN Security Council Consolidated List
    sanctioned = [p.party_id for p in ubos if p.sanctions_flag]
    if UN in kinds or sanctioned:
        add("RF-03", f"UBO 确认命中制裁名单：{', '.join(sanctioned) or '见 UN 名单命中'}；名单：{declared}")

    # RF-04 个人风险：Holds prominent public function（数据中 pep_status = Confirmed）
    peps = [p.party_id for p in ubos if p.pep_status in ("Confirmed", "PEP")]
    if peps:
        add("RF-04", f"UBO 确认为 PEP：{', '.join(peps)}")
    else:
        # RF-05 个人风险：Spouse, partner, child, parent, or close associate
        assoc = [p.party_id for p in ubos if "associate" in p.pep_status.lower()]
        if assoc:
            add("RF-05", f"UBO 为 PEP 密切关联人：{', '.join(assoc)}")

    # RF-06 股权风险：Complex nested structure exceeding 3 intermediate entities
    layers = case.ownership_layers()
    if layers >= 4:
        add("RF-06", f"股权链共 {layers} 层（>= 4）")

    # RF-07 股权风险：Indirect shareholding via nominee or trust
    nominees = [p.party_id for p in case.all_parties() if p.is_nominee or p.is_trust]
    if nominees:
        add("RF-07", f"存在代持/信托安排：{', '.join(nominees)}")

    # RF-08 业务风险：High-risk industry as main business
    # 案例数据以 industry_tier 表达行业风险（I3 = High-Risk / Restricted Industries）
    iname = case.industry.industry_name or ""
    if case.industry.industry_tier == "I3" or any(k in iname.lower() for k in HIGH_RISK_INDUSTRY):
        add("RF-08", f"主业为高风险行业：{iname}（tier {case.industry.industry_tier}）")

    # RF-09 文件风险：Missing director declaration, shareholder register, etc.
    missing = [d.document_id for d in case.documents if d.is_mandatory and d.status == "Missing"]
    if missing:
        add("RF-09", f"必备文件缺失：{', '.join(missing)}")

    # RF-10 文件风险：文件之间就同一事实相互矛盾
    conflict = [d.document_id for d in case.documents if d.status == "Conflicting"]
    if conflict:
        add("RF-10", f"文件冲突：{', '.join(conflict)}")

    return hits


# ======================================================================================
# 5. Stage 2：计分与分带 -> RiskAssessmentResult
# ======================================================================================


def band_of(total: int, rb: RuleBook) -> str:
    for b in rb.bands:
        if total < b.low:
            continue
        if b.high is not None and total > b.high:
            continue
        return b.band
    return rb.bands[-1].band


def status_of(hits: list[RiskFactorHit], hard_stop: bool) -> str:
    fids = {h.factor_id for h in hits}
    if hard_stop:
        return "Declined"
    if "RF-09" in fids:
        return "ReturnedForSupplementation"
    if "RF-10" in fids:
        return "EscalatedToHumanReview"
    return "Scored"


def assess(case: Case, rb: RuleBook) -> tuple[RiskAssessmentResult, list[RiskFactorHit]]:
    hits = detect(case, rb)
    hard_stop = any(h.has_hard_stop for h in hits)
    total = sum(h.points_applied for h in hits if h.action_type == SCORE)

    result = RiskAssessmentResult(
        assessment_id=f"RA-{case.application.application_id}",
        total_score=total,
        risk_band=rb.bands[-1].band if hard_stop else band_of(total, rb),
        has_hard_stop=hard_stop,
        assessment_date=case.assessment_date or case.application.submission_date,
        assessment_status=status_of(hits, hard_stop),
        edd_required=False,  # EDD 不在本次本体范围
    )
    return result, hits


# ======================================================================================
# 6. 案例加载：test_cases_data.json
# ======================================================================================


def _party(d: dict) -> Party:
    return Party(
        party_id=d.get("party_id", ""),
        person_id=d.get("person_id", ""),
        full_name=d.get("full_name", ""),
        is_ubo=bool(d.get("is_ubo", False)),
        is_nominee=bool(d.get("is_nominee", False)),
        is_trust=bool(d.get("is_trust", False)),
        sanctions_flag=bool(d.get("sanctions_flag", False)),
        shareholding_percentage=d.get("shareholding_percentage", 0.0),
        pep_status=d.get("pep_status", "None"),
        nationality=d.get("nationality", ""),
        registration_jurisdiction=d.get("registration_jurisdiction", ""),
    )


def case_from_json(item: dict) -> Case:
    data = item["data"]
    app = data["Application"]
    ind = data["Industry"]
    return Case(
        name=f"{item.get('test_case_id', '')} {item.get('company', '')}".strip(),
        application=Application(app["application_id"], app["submission_date"],
                                app.get("application_status", "")),
        industry=Industry(ind["industry_id"], ind["industry_name"], ind["industry_tier"]),
        party1=_party(data["party1"]),
        party2_chain=[_party(p) for p in data.get("party2", [])],
        documents=[Document(d["document_id"], d["document_type"], d["status"],
                            d["fact_verified"], bool(d.get("is_mandatory", False)))
                   for d in data.get("Document", [])],
        watchlists=[WatchList(w["watchlist_id"], w["list_name"], w.get("list_version_date", ""))
                    for w in data.get("WatchList", [])],
        assessment_date=app["submission_date"],
    )


def load_cases(path: Path) -> list[Case]:
    raw = json.loads(path.read_text(encoding="utf-8"))
    items = raw["test_cases"] if isinstance(raw, dict) else raw
    return [case_from_json(it) for it in items]


# ======================================================================================
# 7. 内置案例（不依赖 json，用于快速自测）
# ======================================================================================

ASOF = "2026-10-06"


def builtin_cases() -> list[Case]:
    def mk(cid, name, industry, tier="I1", nat="GB", pep="None", san=False,
           chain=0, nominee=False, trust=False, lists=(), docs=None, missing=False, conflict=False):
        p1 = Party(f"{cid}-P1", f"{cid}-PER", f"{name}", is_ubo=False, nationality=nat,
                   registration_jurisdiction=nat, shareholding_percentage=100.0)
        chain_parties = [Party(f"{cid}-P2-{i}", f"{cid}-PER2-{i}", is_ubo=(i == chain),
                               pep_status=pep, sanctions_flag=san, is_nominee=nominee,
                               is_trust=trust, nationality=nat, registration_jurisdiction=nat)
                         for i in range(1, chain + 1)] or [
            Party(f"{cid}-P2-1", f"{cid}-PER2-1", is_ubo=True, pep_status=pep,
                  sanctions_flag=san, is_nominee=nominee, is_trust=trust,
                  nationality=nat, registration_jurisdiction=nat)]
        if docs is None:
            docs = [Document(f"{cid}-D1", "DirectorDeclaration",
                             "Missing" if missing else "Provided", "director_identity", True),
                    Document(f"{cid}-D2", "ShareholderRegister",
                             "Conflicting" if conflict else "Provided", "shareholding_pct", True)]
        return Case(
            name=f"{cid} {name}",
            application=Application(cid, ASOF),
            industry=Industry(f"{cid}-IND", industry, tier),
            party1=p1,
            party2_chain=chain_parties,
            documents=docs,
            watchlists=[WatchList(i, n, ASOF) for i, n in lists],
            assessment_date=ASOF,
        )

    BL = ("WL-FATF-BL-202602", "FATF High-Risk Jurisdictions List (Blacklist)")
    GL = ("WL-FATF-GL-202606", "FATF Increased Monitoring List (Grey List)")
    UNL = ("WL-UN-CONSOLIDATED", "UN Security Council Consolidated List")
    return [
        mk("001", "干净客户", "Standard Risk Industries"),
        mk("002", "灰名单", "Standard Risk Industries", nat="VN", lists=[GL]),
        mk("003", "黑名单硬停", "Standard Risk Industries", nat="IR", lists=[BL]),
        mk("004", "制裁硬停", "Standard Risk Industries", san=True, lists=[UNL]),
        mk("005", "PEP", "Standard Risk Industries", pep="Confirmed"),
        mk("006", "股权链 4 层", "Standard Risk Industries", chain=3),
        mk("007", "代持安排", "Standard Risk Industries", chain=2, nominee=True),
        mk("008", "高风险行业", "High-Risk / Restricted Industries", tier="I3"),
        mk("009", "必备文件缺失", "Standard Risk Industries", missing=True),
        mk("010", "文件冲突", "Standard Risk Industries", conflict=True),
        mk("011", "组合 3+3=6", "Standard Risk Industries", nat="VN", chain=2,
           nominee=True, lists=[GL]),
        mk("012", "组合 5+4=9", "High-Risk / Restricted Industries", tier="I3", pep="Confirmed"),
        mk("013", "组合 3+3+4=10", "High-Risk / Restricted Industries", tier="I3",
           nat="VN", chain=2, nominee=True, lists=[GL]),
        mk("014", "硬停与高分并存", "Standard Risk Industries", nat="IR", pep="Confirmed", lists=[BL]),
        mk("015", "PEP 关联人", "Standard Risk Industries", pep="CloseAssociate"),
    ]


# ======================================================================================
# 8. 输出：只打印 RiskAssessmentResult 与 RiskFactorHit
# ======================================================================================


def print_detail(case: Case, res: RiskAssessmentResult, hits: list[RiskFactorHit]) -> None:
    print("=" * 96)
    print(f"{case.name}")
    print(f"  行业 {case.industry.industry_name}（tier {case.industry.industry_tier}）| "
          f"注册地 {case.party1.registration_jurisdiction} | "
          f"UBO {', '.join(p.party_id for p in case.ubos()) or '无'} | "
          f"股权链 {case.ownership_layers()} 层 | 文件 {len(case.documents)} 份 | "
          f"名单 {', '.join(w.watchlist_id for w in case.watchlists) or '无'}")
    print("-" * 96)
    print("RiskAssessmentResult")
    for k in ("assessment_id", "total_score", "risk_band", "has_hard_stop",
              "assessment_date", "assessment_status", "edd_required"):
        print(f"    {k:<18} {getattr(res, k)}")
    print("RiskFactorHit")
    if not hits:
        print("    （无命中）")
    else:
        print(f"    {'hit_id':<9}{'factor_id':<11}{'points_applied':<16}"
              f"{'action_type':<26}{'has_hard_stop':<15}{'edd_triggered'}")
        for h in hits:
            print(f"    {h.hit_id:<9}{h.factor_id:<11}{h.points_applied:<16}"
                  f"{h.action_type:<26}{str(h.has_hard_stop):<15}{h.edd_triggered}")
        print("    证据")
        for h in hits:
            print(f"      {h.factor_id}  {h.evidence}")
    print()


def print_summary(rows) -> None:
    print(f"{'案例':<22}{'命中因子':<34}{'分':>4}  {'带':<8}{'硬停':<7}{'状态'}")
    print("-" * 96)
    for case, res, hits in rows:
        fids = ",".join(h.factor_id for h in hits) or "-"
        print(f"{case.name[:21]:<22}{fids[:33]:<34}{res.total_score:>4}  "
              f"{res.risk_band:<8}{str(res.has_hard_stop):<7}{res.assessment_status}")
    print()


def main() -> int:
    ap = argparse.ArgumentParser(description="IS4920 本体 PoC")
    ap.add_argument("--json", type=Path, default=DEFAULT_JSON, help="测试案例文件")
    ap.add_argument("--builtin", action="store_true", help="改用内置 15 个案例")
    ap.add_argument("--detail", nargs="?", const="all", default=None,
                    help="打印完整结果表（可指定案例号，如 011 / case11）")
    args = ap.parse_args()

    rb = load_rules()
    print(f"规则目录（来自 SSOT）：{len(rb.factors)} 条风险因子，{len(rb.bands)} 个风险带")
    print("  分值：" + " | ".join(
        f"{f.factor_id} {f.action_type}={f.points:+d}" if f.action_type == SCORE
        else f"{f.factor_id} {f.action_type}" for f in rb.factors.values()))
    print("  风险带：" + " | ".join(
        f"{b.band} {b.low}-{b.high if b.high is not None else '∞'}" for b in rb.bands))
    print("  不做：risk appetite / escalation / mitigation / compliance decision / EDD")

    if args.builtin:
        cases = builtin_cases()
        print(f"\n内置案例：{len(cases)} 个\n")
    else:
        cases = load_cases(args.json)
        print(f"\n案例来源：{args.json}（{len(cases)} 个）\n")

    rows = []
    for case in cases:
        res, hits = assess(case, rb)
        rows.append((case, res, hits))

    if args.detail:
        for case, res, hits in rows:
            if args.detail == "all" or args.detail in case.name:
                print_detail(case, res, hits)
    else:
        print_summary(rows)
        print("（加 --detail 查看每个案例的完整结果表，--detail case11 只看单个）")

    n_hs = sum(1 for _, r, _ in rows if r.has_hard_stop)
    print(f"\n合计 {len(rows)} 个案例：硬停 {n_hs} 个，"
          f"{'Band 1 ' + str(sum(1 for _, r, _ in rows if r.risk_band == 'Band 1'))} / "
          f"Band 2 {sum(1 for _, r, _ in rows if r.risk_band == 'Band 2')} / "
          f"Band 3 {sum(1 for _, r, _ in rows if r.risk_band == 'Band 3')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
