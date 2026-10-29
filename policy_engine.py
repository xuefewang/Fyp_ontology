"""Evidence-first implementation of the three IS4920 rule workbooks.

Workbook cells are policy data, never executable instructions. Missing facts are
unknown, not false. Review conclusions must identify their supporting evidence.
"""
from __future__ import annotations

import datetime as dt
import copy
import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

from customer_ontology import inspect_customer, approval_binding_issues, build_report_payload, digest, supported_rule_pack
from ownership import analyze_ownership

LEVELS = {'Low': 0, 'Medium': 1, 'High': 2}
PENDING = 'Pending'


@dataclass
class Rule:
    rule_id: str
    description: str
    condition: str
    outcome: str
    source: str
    points: int = 0
    kind: str = ''
    priority: int = 0
    evidence: str = ''


@dataclass
class Catalogue:
    rf: dict[str, Rule]
    dimensions: dict[str, Rule]
    appetite: dict[str, Rule]
    safeguards: dict[str, Rule]
    integration: dict[str, Rule]
    bands: list[tuple[str, int, int | None]]
    versions: dict[str, str]
    warnings: list[str] = field(default_factory=list)


def load_catalogue(directory: Path) -> Catalogue:
    import openpyxl
    names = {
        'rf': 'IS4920_Integrated_HardGate_RiskFactors_RiskBand.xlsx',
        'aml': 'IS4920_AML_Appetite_Safeguards.xlsx',
        'pi': 'IS4920_Policy_Integration_Final_Step.xlsx',
    }
    books, versions = {}, {}
    for key, name in names.items():
        path = directory / name
        versions[name] = hashlib.sha256(path.read_bytes()).hexdigest()
        books[key] = openpyxl.load_workbook(path, data_only=True)
    def rows(key, sheet, prefix):
        return [r for r in books[key][sheet].iter_rows(values_only=True)
                if isinstance(r[0], str) and re.fullmatch(prefix + r'-\d{2}', r[0])]
    try:
        rf = {}
        for r in rows('rf', 'Risk_Factors', 'RF'):
            if r[6] not in {'HardGate', 'DocGate', 'Score'}:
                raise ValueError(f'Unsupported factor type: {r[0]}')
            if r[6] == 'Score' and (not isinstance(r[4], (int, float)) or r[4] < 0):
                raise ValueError(f'Invalid score: {r[0]}')
            rf[r[0]] = Rule(r[0], r[2], r[3], str(r[4]), names['rf'],
                            int(r[4]) if r[6] == 'Score' else 0, r[6])
        warnings = []
        dims, appetite, sg, pi = {}, {}, {}, {}
        for prefix, sheet in [('D1', 'D1_Ownership'), ('D2', 'D2_Industry'), ('D3', 'D3_Source_of_Funds')]:
            for r in rows('aml', sheet, prefix):
                dims[r[0]] = Rule(r[0], r[1], r[2], r[5].split('\n')[0], names['aml'], evidence=r[3])
        for r in rows('aml', 'Final_Appetite', 'FA'):
            appetite[r[0]] = Rule(r[0], r[1], r[2], r[5].split('\n')[0], names['aml'], evidence=r[3])
        for r in rows('aml', 'Safeguards', 'SG'):
            sg[r[0]] = Rule(r[0], r[1], r[4], r[7], names['aml'], evidence=r[6])
        for r in rows('pi', 'Policy_Integration', 'PI'):
            pi[r[0]] = Rule(r[0], r[2], r[2], r[5].split('\n')[0], names['pi'], priority=int(r[1]), evidence=r[3])
        bands = []
        for r in books['rf']['Risk_Band'].iter_rows(values_only=True):
            if isinstance(r[0], str) and r[0].split(' / ')[0] in LEVELS:
                band = r[0].split(' / ')[0]
                nums = [int(n) for n in re.findall(r'\d+', r[1])]
                bands.append((band, nums[0], nums[1] if len(nums) == 2 else None))
        expected = {f'RF-{i:02}' for i in range(1, 22)}
        if not expected <= rf.keys() or len(dims) != 17 or len(appetite) != 7 or len(sg) != 13 or len(pi) != 11:
            raise ValueError('Workbook rule set is incomplete or incompatible with this engine')
        order = ['PI-01', 'PI-03', 'PI-02', 'PI-11', 'PI-06', 'PI-04', 'PI-05', 'PI-07', 'PI-08', 'PI-09', 'PI-10']
        if [r.rule_id for r in sorted(pi.values(), key=lambda r: r.priority)] != order:
            raise ValueError('Policy priorities do not match the supported final decision order')
        if bands != [('Low', 0, 4), ('Medium', 5, 9), ('High', 10, None)]:
            raise ValueError('Unexpected risk-band definitions; review engine compatibility')
        return Catalogue(rf, dims, appetite, sg, pi, bands, versions, warnings)
    finally:
        for book in books.values():
            book.close()


def yes(facts, *keys):
    return all(facts.get(key) is True for key in keys)


def no(facts, *keys):
    return all(facts.get(key) is False for key in keys)


def refs(record):
    value = record.get('evidence_refs', [])
    return isinstance(value, list) and bool(value) and all(isinstance(x, str) and x.strip() for x in value)


def dimension_hits(prefix, record):
    """Facts represent documented reviewer conclusions, not guessed thresholds."""
    f = record.get('facts', {})
    if not isinstance(f, dict) or not yes(record, 'evidence_complete') or not refs(record):
        return ['D1-07'] if prefix == 'D1' else ['D2-05'] if prefix == 'D2' else ['D3-04'], PENDING
    hits = []
    if prefix == 'D1':
        indicators = ('complex_ownership', 'nominee_or_trust', 'cross_entity_or_border_complexity',
                      'control_rights_discrepancy', 'unusual_structure', 'material_unresolved_contradiction')
        if any(type(f.get(k)) is not bool for k in indicators):
            return [], PENDING
        if not yes(f, 'legal_entity_verified', 'ownership_links_verified', 'ubo_verified', 'control_verified'):
            return ['D1-07'], PENDING
        required_details = []
        if f['complex_ownership'] or f['nominee_or_trust'] or f['cross_entity_or_border_complexity'] or f['unusual_structure']:
            required_details.append('rationale_sufficient')
        if f['control_rights_discrepancy']: required_details.append('discrepancy_explained')
        if any(type(f.get(k)) is not bool for k in required_details):
            return [], PENDING
        depth = f.get('intermediate_legal_entity_count')
        if type(depth) is not int or depth < 0 or f['complex_ownership'] != (depth >= 4):
            return [], PENDING
        if yes(f, 'complex_ownership', 'rationale_sufficient') and no(f, 'material_unresolved_contradiction'):
            hits.append('D1-02')
        if yes(f, 'nominee_or_trust', 'relevant_parties_verified', 'legal_documents_sufficient', 'rationale_sufficient'):
            hits.append('D1-03')
        if yes(f, 'cross_entity_or_border_complexity', 'materiality_policy_approved', 'rationale_sufficient') and no(f, 'material_unresolved_contradiction', 'control_rights_discrepancy', 'unusual_structure'):
            hits.append('D1-04')
        if yes(f, 'control_rights_discrepancy') and no(f, 'discrepancy_explained'):
            hits.append('D1-05')
        if yes(f, 'unusual_structure', 'review_completed') and no(f, 'rationale_sufficient'):
            hits.append('D1-06')
        if no(f, 'complex_ownership', 'nominee_or_trust', 'cross_entity_or_border_complexity', 'control_rights_discrepancy', 'unusual_structure', 'material_unresolved_contradiction'):
            hits.append('D1-01')
    elif prefix == 'D2':
        if any(type(f.get(k)) is not bool for k in ('regulated', 'material_discrepancy', 'material_high_risk_activity')):
            return ['D2-05'], PENDING
        if yes(f, 'material_discrepancy') or not yes(f, 'activity_verified', 'taxonomy_approved'):
            return ['D2-05'], PENDING
        licensed = f.get('regulated') is False or yes(f, 'licence_verified')
        if f.get('category') == 'Standard' and no(f, 'material_discrepancy', 'material_high_risk_activity') and licensed:
            hits.append('D2-01')
        if f.get('category') == 'Sensitive' and licensed and no(f, 'material_high_risk_activity'):
            hits.append('D2-02')
        if f.get('category') == 'HighRisk' and yes(f, 'activity_not_prohibited') and licensed:
            hits.append('D2-03')
        if yes(f, 'multiple_activities', 'material_high_risk_activity', 'materiality_policy_approved', 'high_risk_activity_verified', 'activity_not_prohibited') and licensed:
            hits.append('D2-04')
    else:
        if not yes(f, 'mandatory_evidence_authenticated') or any(type(f.get(k)) is not bool for k in ('material_anomaly', 'unresolved_discrepancy', 'complex_funding', 'unusual_subscription')):
            return ['D3-04'], PENDING
        if f['material_anomaly'] and (type(f.get('anomaly_explained')) is not bool or not yes(f, 'records_verified', 'review_completed')):
            return [], PENDING
        if yes(f, 'material_anomaly', 'records_verified', 'review_completed') and no(f, 'anomaly_explained'):
            hits.append('D3-03')
        reconciled = yes(f, 'source_event_identified', 'owner_identified', 'amount_dates_verified', 'funds_sufficient', 'records_reconciled') and no(f, 'unresolved_discrepancy')
        if reconciled and yes(f, 'own_account', 'records_verified') and no(f, 'complex_funding', 'unusual_subscription', 'material_anomaly'):
            hits.append('D3-01')
        if reconciled and yes(f, 'complex_funding', 'material_parties_verified', 'legal_basis_verified', 'flows_verified', 'arrangement_policy_permitted'):
            hits.append('D3-02')
        if reconciled and yes(f, 'unusual_subscription', 'materiality_policy_approved', 'exceptional_event_verified'):
            hits.append('D3-05')
    return hits, None if hits else PENDING


def evaluate_dimensions(evidence, cat, rf=None):
    results = {}
    for prefix in ('D1', 'D2', 'D3'):
        record = evidence.get(prefix, {})
        hits, status = dimension_hits(prefix, record)
        level = PENDING if status == PENDING else max((cat.dimensions[x].outcome for x in hits), key=LEVELS.__getitem__)
        results[prefix] = {'level': level, 'rule_ids': hits, 'evidence_refs': record.get('evidence_refs', []),
                           'remediation_rf_ids': ['RF-11'] if 'D1-07' in hits else ['RF-09', 'RF-10'] if 'D2-05' in hits else ['RF-09'] if 'D3-04' in hits else []}
    levels = [d['level'] for d in results.values()]
    if PENDING in levels:
        appetite, fid = PENDING, 'FA-04'
    else:
        highest = max(levels, key=LEVELS.__getitem__)
        appetite, fid = {'Low': ('Within', 'FA-01'), 'Medium': ('Conditional', 'FA-02'), 'High': ('Restricted', 'FA-03')}[highest]
    rules = [fid, 'FA-05']
    if (any(d['level'] == PENDING and not d['rule_ids'] for d in results.values())
            or (results['D2']['level'] == PENDING and evidence.get('D2', {}).get('facts', {}).get('taxonomy_approved') is not True)):
        rules.append('FA-07')
    pep_hits = {h['factor_id'] for h in (rf or {}).get('hits', [])} & {'RF-13', 'RF-14'}
    if appetite != PENDING and pep_hits:
        appetite = 'Conditional' if appetite == 'Within' else appetite
        rules.append('FA-06')
    return results, appetite, rules


RF_CRITERIA = {
    'RF-01': ('all', ('prohibited_jurisdiction', 'official_verification_complete')),
    'RF-02': ('all', ('ubo_sanctions_match', 'screening_review_complete', 'false_positive_excluded')),
    'RF-03': ('all', ('regulated_activity', 'licence_missing_revoked_or_invalid', 'official_verification_complete', 'policy_prohibits_onboarding')),
    'RF-04': ('failed', ('ubo_verification_status',)),
    'RF-05': ('failed', ('control_verification_status',)),
    'RF-06': ('all', ('prohibited_or_illegal_activity', 'reliable_evidence_verified', 'approved_prohibition_met')),
    'RF-07': ('all', ('illegal_subscription_funds', 'reliable_evidence_verified', 'approved_prohibition_met')),
    'RF-08': ('all', ('fraudulent_identity_or_key_documents', 'independent_evidence_verified', 'approved_prohibition_met')),
    'RF-09': ('all', ('mandatory_document_missing',)),
    'RF-10': ('all', ('articles_shareholder_register_shareholding_conflict',)),
    'RF-11': ('ubo_pending', ('ubo_verification_status', 'ownership_evidence_insufficient', 'remediation_outstanding')),
    'RF-12': ('all', ('registration_or_ubo_nationality_greylisted',)),
    'RF-13': ('all', ('ubo_confirmed_pep',)),
    'RF-14': ('all', ('ubo_pep_family_or_close_associate',)),
    'RF-15': ('layers', ('intermediate_legal_entity_count',)),
    'RF-16': ('all', ('nominee_or_trust_indirect_shareholding',)),
    'RF-17': ('all', ('principal_business_in_approved_high_risk_category',)),
    'RF-18': ('any', ('third_party_funding', 'non_own_name_account', 'borrowed_or_leveraged_funding')),
    'RF-19': ('all', ('incorporated_under_one_year', 'no_substantive_history_or_audited_accounts')),
    'RF-20': ('any', ('policy_defined_high_leverage', 'negative_cash_flow', 'assets_or_cash_insufficient_for_subscription')),
    'RF-21': ('jurisdictions', ('jurisdiction_count', 'complex_cross_border_or_offshore_arrangement')),
}


def rf_condition(fid, facts):
    """Three-valued predicates: incomplete facts never imply Clear."""
    method, keys = RF_CRITERIA[fid]
    values = [facts.get(k) for k in keys]
    if method == 'failed':
        return True if values[0] == 'Failed' else False if values[0] == 'Verified' else None
    if method == 'ubo_pending':
        if values[0] == 'Failed': return False
        if values[0] in {'Pending', 'AwaitingInformation'} or any(v is True for v in values[1:]): return True
        return False if values[0] == 'Verified' and all(v is False for v in values[1:]) else None
    if method == 'layers':
        depth = values[0]
        return depth >= 4 if type(depth) is int and depth >= 0 else None
    if method == 'jurisdictions':
        count, complex_structure = values
        if (type(count) is int and count >= 3) or complex_structure is True:
            return True
        return False if type(count) is int and count >= 1 and complex_structure is False else None
    if method == 'all':
        if any(v is False for v in values): return False
        return True if all(v is True for v in values) else None
    if any(v is True for v in values): return True
    return False if all(v is False for v in values) else None


def evaluate_rf(evidence, cat):
    records = evidence.get('risk_factors', {})
    unknown = set(records) - set(cat.rf)
    input_errors = [f'Unknown risk factor IDs: {sorted(unknown)}'] if unknown else []
    hits, pending, assessments = [], [], []
    hard_pending, doc_pending = [], []
    for fid, rule in cat.rf.items():
        record = records.get(fid, {})
        if not isinstance(record, dict):
            input_errors.append(f'{fid}: invalid assessment record'); record = {}
        if 'facts' in record:
            condition = rf_condition(fid, record['facts']) if isinstance(record['facts'], dict) else None
            if not isinstance(record['facts'], dict): input_errors.append(f'{fid}: invalid fact record')
            status = 'Hit' if condition is True else 'Clear' if condition is False else PENDING
            supplied = record.get('status')
            if supplied is not None and supplied != status:
                input_errors.append(f'{fid}: declared status contradicts predicate facts')
                status = PENDING
        else:
            status = record.get('status', PENDING)
        if not isinstance(status, str) or status not in {'Hit', 'Clear', PENDING}:
            input_errors.append(f'{fid}: status must be Hit, Clear or Pending')
            status = PENDING
        if status != PENDING and (not refs(record) or record.get('verified') is not True):
            status = PENDING
        assessments.append({'rule_id': fid, 'status': status, 'evidence_refs': record.get('evidence_refs', [])})
        if status == PENDING:
            pending.append(fid)
            if rule.kind == 'HardGate': hard_pending.append(fid)
            if rule.kind == 'DocGate': doc_pending.append(fid)
        elif status == 'Hit':
            hits.append({'factor_id': fid, 'factor_type': rule.kind, 'points_applied': rule.points,
                         'action_type': rule.outcome, 'evidence_refs': record['evidence_refs']})
    hard = 'Confirmed Hit' if any(h['factor_type'] == 'HardGate' for h in hits) else PENDING if hard_pending else 'Clear'
    doc_hits = {h['factor_id'] for h in hits if h['factor_type'] == 'DocGate'}
    doc = 'Human Review' if 'RF-10' in doc_hits or evidence.get('ownership_review_required') is True else 'Awaiting Information' if doc_hits else PENDING if doc_pending else 'Clear'
    subtotal = sum(h['points_applied'] for h in hits)
    score = subtotal if doc == 'Clear' else None
    band = next(b for b, low, high in cat.bands if subtotal >= low and (high is None or subtotal <= high)) if score is not None else PENDING
    # A confirmed prohibition is reported separately, never converted to High.
    final_band = PENDING if hard != 'Clear' or doc != 'Clear' or pending else band
    return {'hard_gate': hard, 'doc_gate': doc, 'total_score': score, 'score_band': band,
            'risk_band': final_band, 'confirmed_score_subtotal': subtotal,
            'scoring_status': 'Completed' if doc == 'Clear' else 'Held — RF document remediation',
            'doc_gate_reasons': ([a for a in assessments if a['rule_id'] in {'RF-09', 'RF-10', 'RF-11'} and a['status'] != 'Clear']
                                 + ([{'rule_id': 'RF-15', 'status': 'Human Review', 'reason': 'Ownership cycle'}]
                                    if evidence.get('ownership_review_required') is True else [])),
            'hits': hits, 'pending_rule_ids': pending, 'rule_assessments': assessments, 'input_errors': input_errors}


def valid_approval(record, roles=None):
    return (record.get('decision') in {'Approved', 'Rejected'}
            and yes(record, 'authorized') and refs(record)
            and all(isinstance(record.get(k), str) and record[k].strip()
                    for k in ('reviewer', 'authority', 'timestamp', 'rationale'))
            and (roles is None or record.get('role') in roles))


def safeguard_candidates(dimensions, rf, appetite, evidence):
    hits = {h['factor_id'] for h in rf['hits']}
    dims = {r for d in dimensions.values() for r in d['rule_ids']}
    signals = evidence.get('signals', {})
    result = set()
    if (appetite == 'Restricted' or dims & {'D1-02', 'D1-05', 'D1-06', 'D2-03', 'D2-04', 'D3-03'}
            or len([h for h in rf['hits'] if h.get('factor_type') == 'Score']) >= 2
            or yes(signals, 'edd_required')):
        result.add('SG-01')
    if dims & {'D3-02', 'D3-03', 'D3-05', 'D2-03', 'D2-04'} or hits & {'RF-13', 'RF-14', 'RF-17'}: result.add('SG-02')
    if bool(hits & {'RF-13', 'RF-14'}) or yes(signals, 'high_risk_relationship'): result.add('SG-03')
    if appetite == 'Restricted' or bool(hits & {'RF-13', 'RF-14'}) or yes(signals, 'high_risk_relationship'): result.add('SG-04')
    if dims & {'D3-02', 'D3-03'} and yes(signals, 'third_party_or_account_mismatch'): result.add('SG-05')
    if yes(signals, 'exposure_limit_reduces_risk'): result.add('SG-06')
    if yes(signals, 'service_restriction_exposure'): result.add('SG-07')
    if appetite in {'Conditional', 'Restricted'}: result.add('SG-08')
    if bool(hits & {'RF-13', 'RF-14'}) or 'RF-12' in hits or yes(signals, 'rescreening_exposure'): result.add('SG-09')
    if appetite == 'Restricted' or yes(signals, 'high_risk_relationship'): result.add('SG-10')
    if dims & {'D1-02', 'D1-03', 'D1-04'} or hits & {'RF-15', 'RF-16'}: result.add('SG-11')
    if yes(signals, 'regulated_activity') or yes(evidence.get('D2', {}).get('facts', {}), 'regulated'): result.add('SG-12')
    if bool(hits & {'RF-13', 'RF-14'}) or yes(signals, 'dual_approval_required'): result.add('SG-13')
    return result


def evaluate_safeguards(dimensions, rf, appetite, evidence, cat):
    records = evidence.get('safeguards', {})
    unknown = set(records) - set(cat.safeguards)
    input_errors = [f'Unknown safeguard IDs: {sorted(unknown)}'] if unknown else []
    candidates = safeguard_candidates(dimensions, rf, appetite, evidence)
    selected, unresolved = [], []
    for sid in sorted((candidates | set(records)) & set(cat.safeguards)):
        record = records.get(sid, {})
        if not isinstance(record, dict):
            input_errors.append(f'{sid}: invalid safeguard record'); record = {}
        # Applicability requires an approved policy and a documented trigger review.
        if type(record.get('applicable')) is not bool or record.get('policy_approved') is not True or not refs(record):
            unresolved.append(sid)
            continue
        if not record['applicable']:
            continue
        status = record.get('status', PENDING)
        if not isinstance(status, str) or status not in {'Complete', 'Incomplete', PENDING}:
            input_errors.append(f'{sid}: invalid safeguard completion status')
            status = PENDING
        if status == 'Complete' and (not isinstance(record.get('owner'), str) or not record['owner'].strip() or not refs({'evidence_refs': record.get('completion_evidence_refs')})
                                     or not yes(record, 'completion_verified')
                                     or not all(isinstance(record.get(k), str) and record[k].strip() for k in ('verifier', 'verified_at'))):
            status = PENDING
        approvals = [a for a in evidence.get('approvals', []) if isinstance(a, dict)
                     and a.get('assessment_id') == evidence.get('assessment_id')
                     and a.get('policy_version') == evidence.get('policy', {}).get('version')]
        if sid == 'SG-04' and status == 'Complete' and not any(valid_approval(a, {'Senior Compliance'}) and a['decision'] == 'Approved' for a in approvals):
            status = PENDING
        if sid == 'SG-13' and status == 'Complete':
            directors = {a['role']: a['reviewer'] for a in approvals
                         if valid_approval(a, {'Compliance Director', 'Risk Director'}) and a['decision'] == 'Approved'}
            if len(directors) != 2 or len(set(directors.values())) != 2:
                status = PENDING
        if sid in {'SG-08', 'SG-09', 'SG-10'} and record.get('status') == 'Complete':
            schedule = record.get('ongoing_setup', {})
            valid_schedule = (isinstance(schedule, dict) and yes(schedule, 'cadence_approved', 'setup_verified', 'signed_off')
                              and all(isinstance(schedule.get(k), str) and schedule[k].strip()
                                      for k in ('cadence', 'first_due_date', 'executable_plan', 'verifier', 'verified_at'))
                              and refs(schedule))
            if not isinstance(schedule, dict) or schedule.get('cadence_approved') is not True:
                status = PENDING
            elif not valid_schedule or not record.get('owner'):
                status = 'Incomplete'
        selected.append({'safeguard_id': sid, 'status': status, 'owner': record.get('owner'),
                         'evidence_refs': record['evidence_refs'],
                         'completion_evidence_refs': record.get('completion_evidence_refs', []),
                         'ongoing_setup': record.get('ongoing_setup') if sid in {'SG-08', 'SG-09', 'SG-10'} else None})
    if not yes(evidence.get('safeguard_selection', {}), 'reviewed') or not refs(evidence.get('safeguard_selection', {})):
        status = PENDING
    elif unresolved or any(x['status'] == PENDING for x in selected):
        status = PENDING
    elif any(x['status'] == 'Incomplete' for x in selected):
        status = 'Incomplete'
    elif selected:
        status = 'Complete'
    elif appetite == PENDING or not (evidence.get('safeguard_selection', {}).get('reviewed') is True
                                     and refs(evidence.get('safeguard_selection', {}))):
        status = PENDING
    else:
        status = 'Not Applicable'
    return {'status': status, 'selected': selected, 'unresolved_applicability': unresolved,
            'supported_rule_ids': sorted(cat.safeguards), 'input_errors': input_errors}


def integrate(rf, appetite, sg, evidence, cat):
    """Consume upstream outcomes in the exact workbook priority order."""
    policy = evidence.get('policy', {})
    required = policy.get('required_approver_roles')
    if isinstance(required, list) and all(isinstance(r, str) for r in required):
        required = list(required)
        for item in sg.get('selected', []):
            roles = {'SG-04': ['Senior Compliance'], 'SG-13': ['Compliance Director', 'Risk Director']}.get(item['safeguard_id'], [])
            required.extend(r for r in roles if r not in required)
    assessment_id = evidence.get('assessment_id')
    version = policy.get('version')
    policy_defined = (yes(policy, 'approval_matrix_approved') and isinstance(version, str) and bool(version.strip())
                      and isinstance(required, list) and bool(required)
                      and all(isinstance(r, str) and r.strip() for r in required)
                      and len(set(required)) == len(required) and bool(assessment_id))
    errors = list(evidence.get('input_errors', [])) + rf.get('input_errors', []) + sg.get('input_errors', [])
    statuses_valid = (rf.get('hard_gate') in {'Confirmed Hit', 'Clear', PENDING}
        and rf.get('doc_gate') in {'Clear', 'Awaiting Information', 'Human Review', PENDING}
        and rf.get('risk_band') in {*LEVELS, PENDING}
        and appetite in {'Within', 'Conditional', 'Restricted', PENDING}
        and sg.get('status') in {'Not Applicable', 'Complete', 'Incomplete', PENDING})
    if not statuses_valid: errors.append('Missing or invalid required upstream status')
    if not policy_defined: errors.append('Undefined approved authority matrix, policy version or assessment identity')
    for name, context in evidence.get('upstream_contexts', {}).items():
        if not isinstance(context, dict) or context.get('assessment_id') != assessment_id or context.get('policy_version') != version:
            errors.append(f'{name}: mismatched assessment or policy version')
    valid = []
    for record in evidence.get('approvals', []):
        if not isinstance(record, dict): errors.append('Invalid approval record'); continue
        if record.get('role') not in (required or []): continue
        if record.get('assessment_id') != assessment_id or record.get('policy_version') != version:
            errors.append('Required approval belongs to a different assessment/policy version'); continue
        if record.get('authorized') is not True:
            errors.append('Required approver authority is unknown or invalid'); continue
        if valid_approval(record): valid.append(record)
    rejected = any(a['decision'] == 'Rejected' for a in valid)
    approved = policy_defined and all(any(a['role'] == r and a['decision'] == 'Approved' for a in valid) for r in required)
    authorization = 'Rejected' if rejected else 'Approved' if approved else PENDING
    conditions = {
        'PI-01': rf.get('hard_gate') == 'Confirmed Hit',
        'PI-03': rf.get('doc_gate') in {'Awaiting Information', 'Human Review', PENDING},
        'PI-02': rf.get('hard_gate') == PENDING,
        'PI-11': bool(errors),
        'PI-06': rejected,
        'PI-04': rf.get('risk_band') == PENDING or appetite == PENDING,
        'PI-05': sg.get('status') in {'Incomplete', PENDING},
        'PI-07': authorization == PENDING,
        'PI-08': appetite == 'Within',
        'PI-09': appetite == 'Conditional',
        'PI-10': appetite == 'Restricted',
    }
    # Shared acceptance guard; Restricted is expressly permitted in this version.
    guard = (rf.get('hard_gate') == 'Clear' and rf.get('doc_gate') == 'Clear'
             and rf.get('risk_band') in LEVELS and appetite in {'Within', 'Conditional', 'Restricted'}
             and sg.get('status') in {'Complete', 'Not Applicable'} and authorization == 'Approved' and not errors)
    for rule in sorted(cat.integration.values(), key=lambda r: r.priority):
        if conditions[rule.rule_id] and (rule.rule_id not in {'PI-08', 'PI-09', 'PI-10'} or guard):
            return {'rule_id': rule.rule_id, 'result': rule.outcome, 'rationale': rule.condition,
                    'authorization_status': authorization, 'required_approver_roles': required, 'input_errors': errors}
    return {'rule_id': 'PI-11', 'result': cat.integration['PI-11'].outcome,
            'rationale': 'No final rule matched the shared acceptance guard',
            'authorization_status': authorization, 'required_approver_roles': required, 'input_errors': errors}


def prepare_workflow_evidence(case, cat_versions=None):
    """Normalize a private evidence copy shared by batch and graph runners."""
    evidence = copy.deepcopy(case.workflow_evidence)
    if not isinstance(evidence, dict):
        evidence = {'input_errors': ['WorkflowEvidence must be an object']}
    for key in ('policy', 'risk_factors', 'safeguards', 'signals', 'D1', 'D2', 'D3', 'upstream_contexts', 'safeguard_selection'):
        if key in evidence and not isinstance(evidence[key], dict):
            evidence.setdefault('input_errors', []).append(f'{key} must be an object')
            evidence[key] = {}
    for prefix in ('D1', 'D2', 'D3'):
        record = evidence.get(prefix, {})
        if 'facts' in record and not isinstance(record['facts'], dict):
            evidence.setdefault('input_errors', []).append(f'{prefix}.facts must be an object')
            record['facts'] = {}
    if 'approvals' in evidence and not isinstance(evidence['approvals'], list):
        evidence.setdefault('input_errors', []).append('approvals must be an array')
        evidence['approvals'] = []
    aid = 'RA-' + case.application.application_id
    if evidence.get('assessment_id') not in (None, aid):
        evidence.setdefault('input_errors', []).append('Input assessment_id does not match current case')
    evidence.setdefault('assessment_id', aid)
    inspection = inspect_customer(case, evidence)
    if inspection['mode'] == 'Customer' and not inspection['issues']:
        relationships = inspection['customer']['party_relationships']
        disclosed = {'investor_id': case.party1.party_id,
                     'nodes': [{'id': p['party_id'], 'entity_type': p['entity_type']} for p in inspection['parties']],
                     'edges': [{'entity_id': r['to_party_id'], 'owner_id': r['from_party_id']}
                               for r in relationships if r['relationship_type'] == 'ownership']}
        if 'ownership_structure' in evidence and evidence['ownership_structure'] != disclosed:
            evidence.setdefault('input_errors', []).append('Ownership relationship graph conflicts with supplied ownership_structure')
        evidence['ownership_structure'] = disclosed
    ownership = None
    if 'ownership_structure' in evidence:
        ownership = analyze_ownership(evidence['ownership_structure'])
        fact = evidence.setdefault('risk_factors', {}).setdefault('RF-15', {})
        declared = fact.pop('status', None)
        count = ownership['intermediate_legal_entity_count']
        if inspection['mode'] == 'Customer':
            d1 = evidence.get('D1', {}).get('facts', {})
            if d1.get('intermediate_legal_entity_count') not in (None, count):
                evidence.setdefault('input_errors', []).append('D1 depth contradicts relationship graph')
        inferred = PENDING if count is None else 'Hit' if count >= 4 else 'Clear'
        if declared is not None and declared != inferred:
            evidence.setdefault('input_errors', []).append('RF-15 conclusion contradicts ownership graph')
        fact['facts'] = {'intermediate_legal_entity_count': count}
        evidence.setdefault('D1', {}).setdefault('facts', {})['intermediate_legal_entity_count'] = ownership['intermediate_legal_entity_count']
        if type(ownership['intermediate_legal_entity_count']) is int:
            evidence['D1']['facts']['complex_ownership'] = ownership['intermediate_legal_entity_count'] >= 4
        if ownership['cycle_detected']:
            evidence['ownership_review_required'] = True
    if inspection['mode'] == 'Customer':
        if cat_versions is None or cat_versions != supported_rule_pack():
            inspection['issues'].append({'code': 'RULE_PACK_COMPATIBILITY', 'path': '$.policy',
                                         'message': 'Rule workbook hashes require explicit engine compatibility review'})
        inspection['issues'].extend(approval_binding_issues(case, evidence, digest(cat_versions))
                                    if cat_versions is not None else [])
        evidence.setdefault('input_errors', []).extend(
            f"{issue['code']}: {issue['message']}" for issue in inspection['issues'])
    evidence['_customer_inspection'] = inspection
    return evidence, ownership


def assess_workflow(case, cat):
    evidence, ownership = prepare_workflow_evidence(case, cat.versions)
    rf = evaluate_rf(evidence, cat)
    dimensions, appetite, appetite_rules = evaluate_dimensions(evidence, cat, rf)
    sg = evaluate_safeguards(dimensions, rf, appetite, evidence, cat)
    final = integrate(rf, appetite, sg, evidence, cat)
    return format_workflow_result(case, cat, evidence, ownership, rf, dimensions, appetite, appetite_rules, sg, final)


def format_workflow_result(case, cat, evidence, ownership, rf, dimensions, appetite, appetite_rules, sg, final):
    """Assemble the same public audit contract for batch and graph execution."""
    ids = [h['factor_id'] for h in rf['hits']] + rf['pending_rule_ids']
    ids += [r for d in dimensions.values() for r in d['rule_ids']] + appetite_rules
    ids += [s['safeguard_id'] for s in sg['selected']] + sg['unresolved_applicability'] + [final['rule_id']]
    result = {'case': case.name, 'assessment_id': 'RA-' + case.application.application_id,
            'assessment_date': case.assessment_date or case.application.submission_date,
            'ownership_analysis': ownership,
            'risk_assessment': rf, 'dimensions': dimensions, 'final_appetite': appetite,
            'appetite_rule_ids': appetite_rules, 'safeguards': sg, 'final_decision': final,
            'edd_required': (True if any(x['safeguard_id'] == 'SG-01' for x in sg['selected'])
                             else None if 'SG-01' in sg['unresolved_applicability'] or appetite == PENDING else False),
            'edd_status': next((x['status'] for x in sg['selected'] if x['safeguard_id'] == 'SG-01'),
                               PENDING if 'SG-01' in sg['unresolved_applicability'] or appetite == PENDING else 'Not Applicable'),
            'audit': {'source_rule_ids': list(dict.fromkeys(ids)), 'workbook_sha256': cat.versions,
                      'policy_version': evidence.get('policy', {}).get('version'),
                      'evaluated_at': dt.datetime.now(dt.timezone.utc).isoformat(),
                      'input_evidence': evidence,
                      'approval_records': evidence.get('approvals', []), 'warnings': cat.warnings}}

    result['report_payload'] = build_report_payload(case, cat, evidence, result)
    return result
